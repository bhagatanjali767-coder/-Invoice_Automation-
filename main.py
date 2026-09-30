from fastapi import FastAPI, UploadFile, File, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from typing import List
from datetime import datetime
import shutil
import os
import re
import pandas as pd
import pytesseract
from PIL import Image
import docx2txt
import docx
import PyPDF2 

# 1. INITIALIZE APP (Only Once)
app = FastAPI(title="Invoice RPA Backend")

# 2. ENABLE CORS MIDDLEWARE
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], 
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 3. ROBUST PATH MANAGEMENT
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FRONTEND_PATH = os.path.join(BASE_DIR, "frontend")
UPLOADS_PATH = os.path.join(BASE_DIR, "uploads")
PROCESSED_PATH = os.path.join(BASE_DIR, "processed_data")

# Ensure required directories exist relative to the workspace root
os.makedirs(UPLOADS_PATH, exist_ok=True)
os.makedirs(PROCESSED_PATH, exist_ok=True)

# 4. MOUNTING STATIC STORAGE
app.mount("/static", StaticFiles(directory=FRONTEND_PATH), name="static")
app.mount("/downloads", StaticFiles(directory=PROCESSED_PATH), name="downloads")

# 5. GLOBAL TESSERACT EXECUTIVE ROUTING
if os.name != 'nt': 
    pytesseract.pytesseract.tesseract_cmd = '/usr/bin/tesseract'
else:
    pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

# ================= EXTRACTION LOGIC =================

def extract_details_from_text(text: str) -> dict:
    """
    Refined extraction logic using advanced Regex.
    Upgraded to handle 15+ different invoice formats and layouts.
    """
    data = {"Invoice Number": "Not Found", "Date": "Not Found", "Total Amount": "Not Found", "Payment Status": "Unpaid"}
    
    # 1. ADVANCED INVOICE NUMBER LOGIC
    inv_keywords = r'(?:Invoice\s*No[\.:]?|Invoice\s*#|Invoice\s*Number|INVOICE\s*#|Challan\s*No|Statement\s*ID|Bill\s*Number|Bill\s*No|Receipt\s*Number|Booking\s*Ref.*?|Job\s*Card\s*No|Order\s*No|Check\s*#|Ref#)'
    inv_pattern = re.search(f'{inv_keywords}\\s*[:\\-]?\\s*\\[?([A-Za-z0-9\\-]+)\\]?', text, re.IGNORECASE)
    
    if inv_pattern:
        data["Invoice Number"] = inv_pattern.group(1)
    else:
        fallback_direct = re.search(r'\b(INV-[A-Z0-9\-]+|GST-[A-Z0-9\-]+)\b', text, re.IGNORECASE)
        if fallback_direct:
            data["Invoice Number"] = fallback_direct.group(1)
        else:
            fallback_txn = re.search(r'(?:Transaction\s*ID|Order\s*ID)\s*[:\-]?\s*([A-Z0-9\-]+)', text, re.IGNORECASE)
            if fallback_txn:
                data["Invoice Number"] = fallback_txn.group(1)

    # 2. ADVANCED DATE LOGIC
    date_patterns = [
        r'\b(\d{1,2}[/\-\.]\d{1,2}[/\-\.]\d{2,4})\b',           # Standard DD/MM/YYYY
        r'\b(\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})\b',               # 22 February 2024
        r'\b([A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4})\b'              # February 22, 2024
    ]
    
    found_date = None
    for pattern in date_patterns:
        match = re.search(pattern, text)
        if match:
            found_date = match.group(1)
            break
            
    if found_date and not any(junk in found_date.lower() for junk in ['pcs', 'qty', 'unit']):
        data["Date"] = found_date
    else:
        data["Date"] = "Not Found"

    # 3. ADVANCED TOTAL AMOUNT LOGIC
    amt_keywords = r'(?:Total\s*Amount|Grand\s*Total|Total\s*Payable|Total\s*Billed|Final\s*Amount|Amount\s*Due|TOTAL|Total\s*Fare|Net\s*Amount|Net\s*Payable|Amount\s*Received|Total\s*Amount\s*After\s*Tax)'
    amount_pattern = re.search(f'{amt_keywords}[\\s:.\\-]*[^\\d]*([0-9,]+\\.[0-9]{{2}})', text, re.IGNORECASE)
    
    if amount_pattern:
        data["Total Amount"] = amount_pattern.group(1)
    else:
        all_amounts = re.findall(r'([0-9]{1,3}(?:,[0-9]{2,3})+\.[0-9]{2}|\d+\.\d{2})', text)
        if all_amounts:
            data["Total Amount"] = all_amounts[-1]

    # 4. ADVANCED PAYMENT STATUS LOGIC
    if re.search(r'\b(?:Paid|Settled|Receipt|Cleared|Successful|Paid in Full|Transaction\s*ID)\b', text, re.IGNORECASE):
        data["Payment Status"] = "Paid"
    elif re.search(r'\b(?:Unpaid|Pending|Due)\b', text, re.IGNORECASE):
        data["Payment Status"] = "Unpaid"

    return data

# ================= APPLICATION ROUTES =================

@app.get("/")
async def serve_frontend():
    """Serves the central UI dashboard index file."""
    return FileResponse(os.path.join(FRONTEND_PATH, "index.html"))

@app.get("/api/status")
async def status():
    """System health check endpoint."""
    return {"status": "online", "message": "Invoice RPA Backend is Online!"}

@app.post("/upload/")
async def upload_invoices(request: Request, files: List[UploadFile] = File(...)):
    extracted_data_list = []

    for file in files:
        file_location = os.path.join(UPLOADS_PATH, file.filename)
        with open(file_location, "wb+") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        extracted_text = ""
        try:
            # 1. IMAGE PROCESSING (OCR)
            if file.filename.lower().endswith(('.png', '.jpg', '.jpeg')):
                img = Image.open(file_location).convert('L')
                width, height = img.size
                img = img.resize((width * 2, height * 2), Image.Resampling.LANCZOS)
                
                # Dynamic config to prevent OS filepath failures on Render (Linux environments)
                if os.name == 'nt':
                    custom_config = r'--tessdata-dir "C:\Program Files\Tesseract-OCR\tessdata" --psm 6'
                else:
                    custom_config = r'--psm 6'
                    
                extracted_text = pytesseract.image_to_string(img, config=custom_config)

            # 2. WORD DOCUMENT PROCESSING
            elif file.filename.lower().endswith('.docx'):
                doc = docx.Document(file_location)
                full_text = []
                for para in doc.paragraphs:
                    if para.text.strip():
                        full_text.append(para.text.strip())
                        
                for table in doc.tables:
                    for row in table.rows:
                        row_data = []
                        for cell in row.cells:
                            clean_text = cell.text.replace('\n', ' ').strip()
                            if clean_text and clean_text not in row_data:
                                row_data.append(clean_text)
                        if row_data:
                            full_text.append("   ".join(row_data))
                            
                extracted_text = "\n".join(full_text)

            # 3. PDF DOCUMENT PROCESSING
            elif file.filename.lower().endswith('.pdf'):
                with open(file_location, 'rb') as pdf_file:
                    pdf_reader = PyPDF2.PdfReader(pdf_file)
                    pdf_text = []
                    for page in pdf_reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            pdf_text.append(page_text)
                    extracted_text = "\n".join(pdf_text)

            print(f"\n--- AI READ THIS FROM: {file.filename} ---")
            print(extracted_text)
            print("-------------------------------------------\n")
            
            parsed_data = extract_details_from_text(extracted_text)
            parsed_data["Source File"] = file.filename
            parsed_data["Processing Status"] = "Success"
            extracted_data_list.append(parsed_data)
            
        except Exception as e:
            extracted_data_list.append({
                "Source File": file.filename, 
                "Invoice Number": "Error",
                "Date": "Error", 
                "Total Amount": "Error", 
                "Payment Status": "Error", 
                "Processing Status": f"Failed: {str(e)}"
            })

    # Prepare for Excel Export
    df = pd.DataFrame(extracted_data_list)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_filename = f"Extracted_Invoices_{timestamp}.xlsx"
    
    excel_output_path = os.path.join(PROCESSED_PATH, excel_filename)
    df.to_excel(excel_output_path, index=False)

    # Dynamic Host Link generation (Works locally and on cloud automatically)
    base_url = str(request.base_url).rstrip("/")
    download_link = f"{base_url}/downloads/{excel_filename}"

    return {
        "status": "success",
        "message": f"Successfully processed {len(files)} invoice(s)!",
        "records_processed": len(files),
        "data_preview": extracted_data_list,
        "download_url": download_link
    }